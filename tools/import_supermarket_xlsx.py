#!/usr/bin/env python3
"""Safe offline supermarket importer. Dry-run by default; never touches other categories."""
import argparse, json, sqlite3, hashlib, shutil
from pathlib import Path
from decimal import Decimal, InvalidOperation
from datetime import datetime, timezone
from openpyxl import load_workbook

def records(path, price_col):
    ws=load_workbook(path,read_only=True,data_only=True).active
    for n,row in enumerate(ws.iter_rows(values_only=True),1):
        values=list(row[:20])+[None]*max(0,20-len(row))
        name=values[1]
        if not isinstance(name,str) or not name.strip(): raise ValueError(f'row {n}: missing name')
        raw=values[price_col-1]
        try: price=Decimal(str(raw))
        except (InvalidOperation,TypeError): raise ValueError(f'row {n}: invalid price in column {price_col}')
        if not price.is_finite() or price<0: raise ValueError(f'row {n}: invalid price')
        yield n,values,name,price

def main():
    p=argparse.ArgumentParser()
    p.add_argument('xlsx');p.add_argument('--db',required=True)
    p.add_argument('--price-column',type=int,required=True,choices=range(1,21),metavar='1..20')
    p.add_argument('--apply',action='store_true');p.add_argument('--merchant-id',type=int)
    a=p.parse_args()
    dbpath=Path(a.db); assert dbpath.is_file(),'Database file not found'
    raw=list(records(a.xlsx,a.price_column))
    if not raw: raise ValueError('empty workbook')
    digest=hashlib.sha256(Path(a.xlsx).read_bytes()).hexdigest()
    print(json.dumps({'rows':len(raw),'sha256':digest,'price_column':a.price_column,'first_three':[(r[2],str(r[3])) for r in raw[:3]],'mode':'APPLY' if a.apply else 'DRY RUN'},ensure_ascii=False))
    if not a.apply: return
    backup=dbpath.with_name(dbpath.name+'.backup-'+datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S'))
    src=sqlite3.connect(dbpath)
    try:
        with sqlite3.connect(backup) as dest: src.backup(dest)
        src.execute('BEGIN IMMEDIATE')
        cols={x[1] for x in src.execute('PRAGMA table_info(products)')}
        if not {'name','category','price','stock','active'}.issubset(cols): raise ValueError('Unexpected products schema')
        if a.merchant_id is not None:
            if 'merchant_id' not in cols: raise ValueError('No merchant support')
            if not src.execute("SELECT 1 FROM merchants WHERE id=? AND category='سوبر ماركت' AND active=1 AND deleted=0",(a.merchant_id,)).fetchone(): raise ValueError('Merchant invalid')
        src.execute('CREATE TABLE IF NOT EXISTS supermarket_import_source (source_sha TEXT NOT NULL, source_row INTEGER NOT NULL, product_id INTEGER NOT NULL, original_columns_json TEXT NOT NULL, PRIMARY KEY(source_sha,source_row))')
        if src.execute('SELECT 1 FROM supermarket_import_source WHERE source_sha=? LIMIT 1',(digest,)).fetchone(): raise ValueError('File already imported')
        src.execute("INSERT OR IGNORE INTO categories(name,active) VALUES ('سوبر ماركت',1)")
        for n,values,name,price in raw:
            fields=['name','category','price','stock','active']; vals=[name,'سوبر ماركت',float(price),100,1]
            for col,val in [('catalog_preview',1),('price_pending',int(price==0)),('merchant_id',a.merchant_id)]:
                if col in cols and (col!='merchant_id' or val is not None): fields.append(col);vals.append(val)
            cur=src.execute('INSERT INTO products('+','.join(fields)+') VALUES ('+','.join('?' for _ in fields)+')',vals)
            src.execute('INSERT INTO supermarket_import_source VALUES (?,?,?,?)',(digest,n,cur.lastrowid,json.dumps(values,ensure_ascii=False,default=str)))
        src.commit()
        print('COMMITTED',len(raw),'backup:',backup)
    except BaseException:
        src.rollback();raise
    finally: src.close()
if __name__=='__main__': main()
